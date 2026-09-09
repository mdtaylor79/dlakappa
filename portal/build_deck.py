#!/usr/bin/env python3
"""
DLA Kappa — Committee Report deck generator
Denton-Lewisville (TX) Alumni Chapter, Kappa Alpha Psi Fraternity, Inc.

Turns committee report files into a PowerPoint deck that matches the chapter's
four-box committee report template — one slide per committee.

Usage
-----
  # every report file in a folder, alphabetical
  python3 build_deck.py reports/ -o chapter-meeting.pptx

  # specific files, in the order you want them presented
  python3 build_deck.py guide-right.json scholarship.json social-action.json

  # a blank template slide to hand out
  python3 build_deck.py --blank -o committee-report-template.pptx

Report file format (one object, or an array of them):
  {
    "committee":    "Guide Right",
    "chair":        "Bro. Chair Name",       # tracked, not printed on the slide
    "meeting":      "2026-10-06",            # tracked, not printed on the slide
    "actions":      "line\nline\nline",      # one bullet per line
    "celebrations": "line\nline",
    "asks":         "line\nline",
    "days":         "Oct 18 - line\nNov 8 - line",
    "titleFont":    "Poppins"                # optional
  }

Requires: python-pptx  (pip install python-pptx)
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# --------------------------------------------------------------------------
# Geometry. The template is authored on a 1280 x 720 px canvas, which maps
# 1:1 onto a 13.333 x 7.5 in widescreen slide at 96 px per inch.
# --------------------------------------------------------------------------
PX = 96.0
SLIDE_W_PX, SLIDE_H_PX = 1280, 720

MARGIN = 22
COL_W, BOX_H = 609, 262
COL_X = (MARGIN, 649)
ROW_Y = (114, 388)

TITLE_TOP, TITLE_H = 34, 56
RULE_TOP, RULE_BOTTOM, RULE_H = 98, 657, 2

ICONS_X, ICONS_Y, ICONS_W, ICONS_H = MARGIN, 667, 140, 47
URL_X, URL_Y, URL_W, URL_H = 190, 676, 900, 30
LOGO_X, LOGO_Y, LOGO_W, LOGO_H = 1198, 660, 59.6, 62

CRIMSON = RGBColor(0x70, 0x11, 0x0C)   # branding guide: chapter crimson
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x9A, 0x9A, 0x9A)      # unfilled placeholder text

TITLE_PT, HEADER_PT, URL_PT = 33, 16.5, 15.75
BODY_PT = {4: 15.75, 5: 14.25, 6: 12.75}   # shrinks as a chair writes more

ASSETS = Path(__file__).parent / "assets"
LOGO_FILE = ASSETS / "dla-logo-small.png"
ICONS_FILE = ASSETS / "social-icons.png"

BOXES = (
    ("actions", "Committee Actions"),
    ("celebrations", "Celebrations"),
    ("asks", "Ask of the Members"),
    ("days", "30 – 60 – 90 Days"),
)
PLACEHOLDER = {
    "actions": ["Enter brief information here"] * 4,
    "celebrations": ["Enter brief information here"] * 4,
    "asks": ["Enter brief information here"] * 4,
    "days": ["Date – Enter the information here"] * 4,
}
BODY_FONT = "Arial"
DEFAULT_TITLE_FONT = "Poppins"


def px(v: float) -> Emu:
    return Inches(v / PX)


def flatten(shape) -> None:
    """Drop the theme style reference so the shape renders flat — no drop shadow,
    no theme tint — in PowerPoint, Keynote, Google Slides and LibreOffice alike."""
    from pptx.oxml.ns import qn

    shape.shadow.inherit = False
    style = shape.element.find(qn("p:style"))
    if style is not None:
        shape.element.remove(style)


def lines(value) -> list[str]:
    if isinstance(value, (list, tuple)):
        items = [str(v).strip() for v in value]
    else:
        items = [s.strip() for s in str(value or "").split("\n")]
    return [s for s in items if s]


# --------------------------------------------------------------------------
# Slide construction
# --------------------------------------------------------------------------
def add_rule(slide, top: int) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, px(MARGIN), px(top), px(SLIDE_W_PX - 2 * MARGIN), px(RULE_H)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = CRIMSON
    bar.line.fill.background()
    flatten(bar)


def add_box(slide, x: int, y: int) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(COL_W), px(BOX_H))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = BLACK
    box.line.width = Pt(1)
    flatten(box)


def text_frame(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def set_run(run, text, *, font, size, bold=False, color=BLACK):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def build_slide(prs: Presentation, report: dict, *, placeholder_color=GREY) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    title_font = (report.get("titleFont") or DEFAULT_TITLE_FONT).strip() or DEFAULT_TITLE_FONT
    name = (report.get("committee") or "").strip() or "(Enter Committee Name)"
    if (report.get("titleCase") or "as") == "upper":
        name = name.upper()

    tf = text_frame(slide, MARGIN, TITLE_TOP, SLIDE_W_PX - 2 * MARGIN, TITLE_H,
                    anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), name, font=title_font, size=TITLE_PT, bold=True)

    add_rule(slide, RULE_TOP)
    add_rule(slide, RULE_BOTTOM)

    for i, (key, label) in enumerate(BOXES):
        x, y = COL_X[i % 2], ROW_Y[i // 2]
        add_box(slide, x, y)

        head = text_frame(slide, x, y + 8, COL_W, 30)
        hp = head.paragraphs[0]
        hp.alignment = PP_ALIGN.CENTER
        set_run(hp.add_run(), label, font=BODY_FONT, size=HEADER_PT, bold=True)

        items = lines(report.get(key))
        placeholder = not items
        if placeholder:
            items = PLACEHOLDER[key]
        size = BODY_PT.get(min(max(len(items), 4), 6), BODY_PT[6])
        color = placeholder_color if placeholder else BLACK

        body = text_frame(slide, x + 14, y + 44, COL_W - 28, BOX_H - 56)
        for n, item in enumerate(items):
            para = body.paragraphs[0] if n == 0 else body.add_paragraph()
            para.alignment = PP_ALIGN.LEFT
            para.space_after = Pt(6)
            set_run(para.add_run(), "•  " + item, font=BODY_FONT, size=size, color=color)

    if ICONS_FILE.exists():
        slide.shapes.add_picture(str(ICONS_FILE), px(ICONS_X), px(ICONS_Y),
                                 px(ICONS_W), px(ICONS_H))

    url = text_frame(slide, URL_X, URL_Y, URL_W, URL_H, anchor=MSO_ANCHOR.MIDDLE)
    up = url.paragraphs[0]
    up.alignment = PP_ALIGN.CENTER
    set_run(up.add_run(), "www.dlakappa.org", font=BODY_FONT, size=URL_PT, bold=True)

    if LOGO_FILE.exists():
        slide.shapes.add_picture(str(LOGO_FILE), px(LOGO_X), px(LOGO_Y),
                                 px(LOGO_W), px(LOGO_H))


def build_deck(reports: list[dict], out: Path, *, placeholder_color=GREY) -> Path:
    prs = Presentation()
    prs.slide_width = px(SLIDE_W_PX)
    prs.slide_height = px(SLIDE_H_PX)
    for report in reports:
        build_slide(prs, report, placeholder_color=placeholder_color)
    props = prs.core_properties
    props.title = "Committee Reports"
    props.author = "Denton-Lewisville (TX) Alumni Chapter"
    props.category = "Kappa Alpha Psi Fraternity, Inc."
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


# --------------------------------------------------------------------------
# Input
# --------------------------------------------------------------------------
def load_reports(paths: list[Path]) -> list[dict]:
    files: list[Path] = []
    for p in paths:
        if p.is_dir():
            files.extend(sorted(p.glob("*.json")))
        elif p.exists():
            files.append(p)
        else:
            print(f"  skipped (not found): {p}", file=sys.stderr)

    reports: list[dict] = []
    for f in files:
        try:
            data = json.loads(f.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, UnicodeDecodeError) as exc:
            print(f"  skipped (not valid JSON): {f.name} — {exc}", file=sys.stderr)
            continue
        batch = data if isinstance(data, list) else [data]
        for item in batch:
            if isinstance(item, dict):
                reports.append(item)
            else:
                print(f"  skipped (not a report): {f.name}", file=sys.stderr)
    return reports


def main() -> int:
    ap = argparse.ArgumentParser(
        description="Build a DLA Kappa committee report deck from report files.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="One slide per report, in the order the files are given.",
    )
    ap.add_argument("inputs", nargs="*", type=Path,
                    help="report .json files and/or folders of them")
    ap.add_argument("-o", "--out", type=Path, default=Path("committee-reports.pptx"),
                    help="output .pptx (default: committee-reports.pptx)")
    ap.add_argument("--blank", action="store_true",
                    help="build one empty template slide instead of reading files")
    args = ap.parse_args()

    if args.blank:
        reports = [{}]
    else:
        if not args.inputs:
            ap.error("give at least one report file or folder, or use --blank")
        reports = load_reports(args.inputs)
        if not reports:
            print("No usable reports found — nothing to build.", file=sys.stderr)
            return 1

    out = build_deck(reports, args.out, placeholder_color=BLACK if args.blank else GREY)
    n = len(reports)
    print(f"{out}  —  {n} slide{'s' if n != 1 else ''}")
    for r in reports:
        print(f"  · {(r.get('committee') or '(untitled)')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
